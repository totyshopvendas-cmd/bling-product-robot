"""Testes do Exportar Planilha JonhDrop — parsing da listagem, mapeamento de
campos da página de detalhe e geração do .xlsx. NÃO acessa JonhDrop/Bling."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "test_export")

import johndrop_export_sheet as sheet  # noqa: E402
import johndrop_export as exp  # noqa: E402

HEADERS = ["Id", "Canal", "Integração", "Img", "Nome", "Sku", "Preço", "Estoque", "Data", "Ações"]


def _row(jd, canal, integ, nome, sku, preco, est, links=None):
    return {
        "cells": [jd, canal, integ, "", nome, sku, preco, est, "01/09/2026 10:00", ""],
        "imgs": ["", "", "", f"https://cdn.x/{jd}.jpg", "", "", "", "", "", ""],
        "links": links or [],
    }


def test_clean_sku_cell_with_inline_badge():
    assert sheet.clean_sku_cell("GDR1015 Catálogo: GDR1015") == "GDR1015"
    assert sheet.clean_sku_cell("KA-966\nCatalogo: KA966") == "KA-966"
    assert sheet.catalog_sku_from_cell("KA-966\nCatálogo: KA966") == "KA966"


def test_parse_brl():
    assert sheet.parse_brl("R$ 1.234,56") == 1234.56
    assert sheet.parse_brl("R$32,50") == 32.5
    assert sheet.parse_brl("21.99") == 21.99
    assert sheet.parse_brl("") is None


def test_parse_list_rows_maps_by_header():
    rows = [_row("119689", "Bling", "TotyShop-Bling", "Copo 260ml", "Copo-260 Catálogo: C260", "R$ 9,99", "30")]
    out = exp.parse_list_rows(HEADERS, rows)
    assert len(out) == 1
    p = out[0]
    assert p["jonhdrop_id"] == "119689"
    assert p["sku"] == "Copo-260"
    assert p["sku_catalogo"] == "C260"
    assert p["preco"] == 9.99
    assert p["estoque"] == 30
    assert p["imagem_principal"] == "https://cdn.x/119689.jpg"


def test_merge_by_sku_groups_integrations():
    rows = [
        _row("1", "Bling", "TotyShop-Bling", "Copo", "C1", "R$ 9,99", "30"),
        _row("2", "Shopee", "TotyShop-Shopee", "Copo", "C1", "R$ 9,99", "30"),
        _row("3", "Bling", "TotyShop-Bling", "Prato", "P1", "R$ 5,00", "3"),
    ]
    merged = exp.merge_by_sku(exp.parse_list_rows(HEADERS, rows))
    assert [m["sku"] for m in merged] == ["C1", "P1"]
    assert merged[0]["canal"] == "Bling, Shopee"
    assert merged[0]["integracao"] == "TotyShop-Bling, TotyShop-Shopee"


def test_pick_detail_link_prefers_edit_with_id():
    item = {
        "jonhdrop_id": "119689",
        "links": [
            "https://app.jonhdrop.com.br/dashboard/product",
            "https://app.jonhdrop.com.br/dashboard/product/delete/119689",
            "https://app.jonhdrop.com.br/dashboard/product/edit/119689",
        ],
    }
    assert exp.pick_detail_link(item).endswith("/edit/119689")
    assert exp.pick_detail_link({"jonhdrop_id": "1", "links": []}) is None


def test_match_field_rules():
    assert sheet.match_field("Preço de Custo (R$)") == "custo"
    assert sheet.match_field("Código de barras") == "gtin"
    assert sheet.match_field("Peso (kg)") == "peso"
    assert sheet.match_field("Peso bruto") == "peso_bruto"
    assert sheet.match_field("Comprimento (cm)") == "comprimento"
    assert sheet.match_field("", "ncm") == "ncm"
    assert sheet.match_field("Descrição do produto") == "descricao"
    assert sheet.match_field("Algo aleatório") is None


def test_map_detail_fields_first_wins():
    raw = [
        {"label": "Custo", "name": "cost", "value": "9,99"},
        {"label": "Custo antigo", "name": "old_cost", "value": "7,00"},
        {"label": "Altura (cm)", "name": "height", "value": "12"},
        {"label": "", "name": "weight", "value": ""},
    ]
    m = sheet.map_detail_fields(raw)
    assert m == {"custo": "9,99", "altura": "12"}


def test_build_workbook(tmp_path):
    from openpyxl import load_workbook

    products = [{
        "jonhdrop_id": "119689", "sku": "Copo-260", "nome": "Copo 260ml com tampa",
        "preco": 9.99, "custo": "R$ 7,50", "estoque": 30,
        "variacoes_lista": ["Rosa", "Azul"], "variacoes_qtd": {"Rosa": 0},
        "variacoes": "Rosa, Azul", "qtd_variacoes": 2,
        "imagens": ["https://cdn.x/a.jpg", "https://cdn.x/b.jpg"], "qtd_imagens": 2,
        "descricao": "=NÃO é fórmula\x01", "link": "https://app.jonhdrop.com.br/dashboard/product/edit/119689",
        "campos_brutos": [{"label": "Altura", "value": "12"}],
    }]
    out = sheet.build_workbook(products, tmp_path / "t.xlsx", {"Origem": "teste"})
    wb = load_workbook(out)
    assert wb.sheetnames == ["Produtos", "Variações", "Imagens", "Todos os campos", "Resumo"]
    ws = wb["Produtos"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["SKU"] == "Copo-260"
    assert row["Custo"] == 7.5
    assert row["Estoque"] == 30
    assert row["Descrição"].startswith("'=")  # protegido contra fórmula
    assert wb["Variações"].max_row == 3
    assert wb["Variações"]["D2"].value == 0
    assert wb["Imagens"].max_row == 3
